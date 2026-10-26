"""Trip-sheet reports: billing (approved sheets only), missing sheets, crew hours, and the printable sheet.

Billing reuses trip_sheet_service.row_billing / summarize_sheet, so a report always agrees with the sheets.
"""
import datetime as dt
import io
from collections import defaultdict

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.models.enums import CrewMode, CrewRole, RowStatus, TripSheetStatus
from app.models.models import BusSchedule, CostCenter, TripReason, TripSheet, TripSheetCrew
from app.services import settings_service
from app.services import trip_sheet_service as ts
from app.services.schedule_service import summarize, version_for_date

MAX_DAYS = 92
ATTRS = ("OPERATOR", "CLIENT", "EXTERNAL")


def check_range(date_from: dt.date, date_to: dt.date):
    if date_to < date_from:
        raise HTTPException(422, "date_to is before date_from")
    if (date_to - date_from).days >= MAX_DAYS:
        raise HTTPException(422, f"Pick a range of at most {MAX_DAYS} days")


def _cc_code(db: Session, cc_id: int) -> str:
    cc = db.get(CostCenter, cc_id)
    return cc.cost_center_code if cc else str(cc_id)


def _bus_label(s: TripSheet) -> str:
    """'EV-101', or 'EV-101 → EV-102' when another bus took over during the day."""
    if not s.bus_change_registration:
        return s.vehicle.fleet_number
    other = s.bus_change_vehicle.fleet_number if s.bus_change_vehicle else s.bus_change_registration
    return f"{s.vehicle.fleet_number} → {other}"


def _r(v) -> float:
    return round(float(v or 0), 2)


# --------------------------------------------------------------------------
# Billing
# --------------------------------------------------------------------------
def billing(db: Session, cc_ids: list[int], date_from: dt.date, date_to: dt.date) -> dict:
    check_range(date_from, date_to)
    sheets = (db.query(TripSheet).filter(TripSheet.cost_center_id.in_(cc_ids), TripSheet.status == TripSheetStatus.APPROVED,
                                         TripSheet.operating_date.between(date_from, date_to))
              .order_by(TripSheet.operating_date, TripSheet.schedule_id).all())
    groups: dict[tuple, dict] = {}
    daily, losses = [], []
    cfg_cache = {}
    for s in sheets:
        cfg = cfg_cache.setdefault(s.cost_center_id, settings_service.get_all(db, s.cost_center_id))
        summ, rb = ts.summarize_sheet(s, cfg)
        run_km = sum(b["run_km"] for r in s.rows for b in [rb[r.row_id]] if r.activity is not None)
        lost_by = {a: 0.0 for a in ATTRS}
        for r in sorted(s.rows, key=lambda x: x.row_sequence):
            b = rb[r.row_id]
            if b["lost_km"] and b["billable"]:
                lost_by[r.attributable_to or "OPERATOR"] += b["lost_km"]
                losses.append({
                    "operating_date": s.operating_date, "cost_center": _cc_code(db, s.cost_center_id),
                    "schedule_code": s.schedule.schedule_code, "fleet_number": s.vehicle.fleet_number,
                    "sheet_id": s.sheet_id, "run": ts._row_label(r), "status": r.status,
                    "reason": r.reason.reason_name if r.reason else None, "attributable_to": r.attributable_to,
                    "planned_km": _r(b["planned_km"]), "run_km": _r(b["run_km"]), "lost_km": _r(b["lost_km"]),
                    "billed": b["lost_billed_km"] > 0})
        day = {"operating_date": s.operating_date, "cost_center_id": s.cost_center_id,
               "cost_center": _cc_code(db, s.cost_center_id), "schedule_id": s.schedule_id,
               "schedule_code": s.schedule.schedule_code, "fleet_number": _bus_label(s),
               "sheet_id": s.sheet_id, "paper_sheet_no": s.paper_sheet_no,
               "planned_billable_km": summ["planned_billable_km"], "lost_operator_km": _r(lost_by["OPERATOR"]),
               "lost_client_km": _r(lost_by["CLIENT"]), "lost_external_km": _r(lost_by["EXTERNAL"]),
               "lost_billed_km": summ["lost_billed_km"], "deducted_km": summ["deducted_km"],
               "billable_km": summ["billable_km"], "odometer_km": summ["odometer_km"], "run_km": _r(run_km),
               "trips_cancelled": summ["trips_cancelled"], "trips_partial": summ["trips_partial"]}
        daily.append(day)
        key = (s.cost_center_id, s.schedule_id)
        g = groups.setdefault(key, {"cost_center_id": s.cost_center_id, "cost_center": day["cost_center"],
                                    "schedule_id": s.schedule_id, "schedule_code": day["schedule_code"],
                                    "fleet_numbers": set(), "days": 0, **{k: 0.0 for k in _SUM_KEYS}})
        g["fleet_numbers"].add(day["fleet_number"])
        g["days"] += 1
        for k in _SUM_KEYS:
            g[k] += day[k]
    rows = []
    for g in groups.values():
        g["fleet_numbers"] = ", ".join(sorted(g["fleet_numbers"]))
        for k in _SUM_KEYS:
            g[k] = _r(g[k])
        g["deviation_pct"] = round((g["odometer_km"] - g["run_km"]) / g["run_km"] * 100, 1) if g["run_km"] and g["odometer_km"] else None
        rows.append(g)
    rows.sort(key=lambda g: (g["cost_center"], g["schedule_code"]))
    totals = {k: _r(sum(g[k] for g in rows)) for k in _SUM_KEYS}
    totals["days"] = sum(g["days"] for g in rows)
    pending = (db.query(TripSheet).filter(TripSheet.cost_center_id.in_(cc_ids), TripSheet.status != TripSheetStatus.APPROVED,
                                          TripSheet.operating_date.between(date_from, date_to)).count())
    return {"date_from": date_from, "date_to": date_to, "rows": rows, "totals": totals, "daily": daily,
            "losses": losses, "not_approved_sheets": pending}


_SUM_KEYS = ("planned_billable_km", "lost_operator_km", "lost_client_km", "lost_external_km", "lost_billed_km",
             "deducted_km", "billable_km", "odometer_km", "run_km", "trips_cancelled", "trips_partial")


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF7")
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    return ws


def billing_xlsx(rep: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Summary", ["Cost center", "Schedule", "Bus", "Days", "Planned billable km", "Lost - operator km",
                           "Lost - client km", "Lost - external km", "Lost but billed km", "Deducted km",
                           "Billable km", "Odometer km", "Deviation %"],
           [[g["cost_center"], g["schedule_code"], g["fleet_numbers"], g["days"], g["planned_billable_km"],
             g["lost_operator_km"], g["lost_client_km"], g["lost_external_km"], g["lost_billed_km"], g["deducted_km"],
             g["billable_km"], g["odometer_km"], g["deviation_pct"]] for g in rep["rows"]]
           + [["TOTAL", "", "", rep["totals"]["days"], *[rep["totals"][k] for k in (
               "planned_billable_km", "lost_operator_km", "lost_client_km", "lost_external_km", "lost_billed_km",
               "deducted_km", "billable_km", "odometer_km")], None]],
           [12, 14, 14, 7, 12, 12, 12, 12, 12, 12, 12, 12, 11])
    _sheet(wb, "Daily", ["Date", "Cost center", "Schedule", "Bus", "Paper sheet", "Planned billable km", "Lost - operator",
                         "Lost - client", "Lost - external", "Lost but billed", "Deducted km", "Billable km",
                         "Odometer km", "Trips cancelled", "Trips cut short"],
           [[d["operating_date"], d["cost_center"], d["schedule_code"], d["fleet_number"], d["paper_sheet_no"],
             d["planned_billable_km"], d["lost_operator_km"], d["lost_client_km"], d["lost_external_km"],
             d["lost_billed_km"], d["deducted_km"], d["billable_km"], d["odometer_km"], d["trips_cancelled"],
             d["trips_partial"]] for d in rep["daily"]],
           [12, 12, 14, 12, 12, 12, 11, 11, 11, 11, 11, 11, 11, 10, 10])
    _sheet(wb, "Losses", ["Date", "Cost center", "Schedule", "Bus", "Run", "Status", "Reason", "Attributable to",
                          "Planned km", "Run km", "Lost km", "Still billed"],
           [[x["operating_date"], x["cost_center"], x["schedule_code"], x["fleet_number"], x["run"], x["status"],
             x["reason"], x["attributable_to"], x["planned_km"], x["run_km"], x["lost_km"], "Yes" if x["billed"] else "No"]
            for x in rep["losses"]],
           [12, 12, 14, 12, 20, 11, 26, 14, 10, 10, 10, 10])
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, dt.date):
                    c.number_format = "DD-MMM-YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Missing sheets
# --------------------------------------------------------------------------
def missing_sheets(db: Session, cc_ids: list[int], date_from: dt.date, date_to: dt.date) -> dict:
    check_range(date_from, date_to)
    schedules = db.query(BusSchedule).filter(BusSchedule.cost_center_id.in_(cc_ids)).order_by(BusSchedule.schedule_code).all()
    sheets = {(s.schedule_id, s.operating_date): s for s in db.query(TripSheet).filter(
        TripSheet.cost_center_id.in_(cc_ids), TripSheet.operating_date.between(date_from, date_to))}
    out, counts, expected = [], defaultdict(int), 0
    day = date_from
    while day <= date_to:
        for sch in schedules:
            v = version_for_date(sch, day)
            if v is None:
                continue
            expected += 1
            sh = sheets.get((sch.schedule_id, day))
            status = sh.status if sh else "NOT_STARTED"
            counts[status] += 1
            if status != TripSheetStatus.APPROVED:
                out.append({"operating_date": day, "cost_center_id": sch.cost_center_id,
                            "cost_center": _cc_code(db, sch.cost_center_id), "schedule_id": sch.schedule_id,
                            "schedule_code": sch.schedule_code, "fleet_number": sch.vehicle.fleet_number,
                            "version_number": v.version_number, "status": status,
                            "sheet_id": sh.sheet_id if sh else None})
        day += dt.timedelta(days=1)
    return {"expected": expected, "counts": dict(counts), "rows": out}


# --------------------------------------------------------------------------
# Crew hours
# --------------------------------------------------------------------------
def _abs(d: dt.date, offset, t):
    return dt.datetime.combine(d, t) + dt.timedelta(days=offset or 0) if t else None


def crew_hours(db: Session, cc_ids: list[int], date_from: dt.date, date_to: dt.date) -> dict:
    check_range(date_from, date_to)
    duties = (db.query(TripSheetCrew).join(TripSheet).filter(
        TripSheet.cost_center_id.in_(cc_ids), TripSheet.operating_date.between(date_from, date_to),
        TripSheetCrew.employee_id.isnot(None)).all())
    items = []
    for c in duties:
        s = c.sheet
        a, b = _abs(s.operating_date, c.duty_start_day_offset, c.duty_start_time), _abs(s.operating_date, c.duty_end_day_offset, c.duty_end_time)
        items.append({"crew_id": c.crew_id, "employee_id": c.employee_id, "employee_number": c.employee.employee_number,
                      "employee_name": c.employee.full_name, "crew_role": c.crew_role, "operating_date": s.operating_date,
                      "cost_center": _cc_code(db, s.cost_center_id), "schedule_code": s.schedule.schedule_code,
                      "fleet_number": s.vehicle.fleet_number, "shift_name": c.shift.shift_name if c.shift else None,
                      "sheet_id": s.sheet_id, "sheet_status": s.status, "start": a, "end": b,
                      "minutes": int((b - a).total_seconds() // 60) if a and b and b > a else 0,
                      "handover_reason": c.handover_reason, "took_over": c.replaces_crew_id is not None,
                      "double_booked": False})
    by_emp = defaultdict(list)
    for it in items:
        by_emp[it["employee_id"]].append(it)
    employees = []
    for emp_id, lst in by_emp.items():
        lst.sort(key=lambda x: (x["start"] or dt.datetime.min))
        for i, x in enumerate(lst):
            for y in lst[i + 1:]:
                if x["start"] and x["end"] and y["start"] and y["start"] < x["end"] and x["start"] < y["end"]:
                    x["double_booked"] = y["double_booked"] = True
        employees.append({"employee_id": emp_id, "employee_number": lst[0]["employee_number"],
                          "employee_name": lst[0]["employee_name"], "days": len({x["operating_date"] for x in lst}),
                          "duties": len(lst), "minutes": sum(x["minutes"] for x in lst),
                          "double_booked": any(x["double_booked"] for x in lst), "items": lst})
    employees.sort(key=lambda e: e["employee_name"])
    return {"date_from": date_from, "date_to": date_to, "employees": employees,
            "total_minutes": sum(e["minutes"] for e in employees)}


def crew_xlsx(rep: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Summary", ["Emp. no.", "Name", "Days", "Duties", "Hours", "Double booked"],
           [[e["employee_number"], e["employee_name"], e["days"], e["duties"], round(e["minutes"] / 60, 2),
             "Yes" if e["double_booked"] else ""] for e in rep["employees"]], [12, 26, 8, 8, 10, 13])
    _sheet(wb, "Duties", ["Date", "Emp. no.", "Name", "Role", "Cost center", "Schedule", "Bus", "Shift", "From", "To",
                          "Hours", "Took over", "Handover reason", "Sheet status", "Double booked"],
           [[x["operating_date"], x["employee_number"], x["employee_name"], x["crew_role"], x["cost_center"],
             x["schedule_code"], x["fleet_number"], x["shift_name"], x["start"], x["end"], round(x["minutes"] / 60, 2),
             "Yes" if x["took_over"] else "", x["handover_reason"], x["sheet_status"], "Yes" if x["double_booked"] else ""]
            for e in rep["employees"] for x in e["items"]],
           [12, 11, 24, 11, 11, 13, 11, 8, 17, 17, 8, 9, 24, 11, 12])
    ws = wb["Duties"]
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "DD-MMM-YYYY"
        row[8].number_format = row[9].number_format = "DD-MMM-YYYY HH:MM"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --------------------------------------------------------------------------
# Printable sheet: the published plan with blanks for the driver
# --------------------------------------------------------------------------
def print_language(cc: CostCenter | None) -> str:
    """The regional language for a cost center's printed sheets.

    The cost center's own choice wins; where it has none - the usual case - the company's is used,
    so a company that prints in Telugu sets it once.
    """
    if cc is None:
        return ""
    if cc.print_language is not None:
        return cc.print_language
    return (cc.company.print_language if cc.company else "") or ""


def print_data(db: Session, schedule: BusSchedule, on: dt.date) -> dict | None:
    v = version_for_date(schedule, on)
    if v is None:
        return None
    cc = schedule.cost_center
    sheet = db.query(TripSheet).filter(TripSheet.schedule_id == schedule.schedule_id, TripSheet.operating_date == on).first()
    roles = [CrewRole.DRIVER] + ([CrewRole.CONDUCTOR] if v.conductor_required else [])
    if v.crew_mode == CrewMode.SHIFTS and v.shifts:
        crew = [{"role": r, "shift": sh.shift_name, "from": sh.start_time.strftime("%H:%M"), "from_day": sh.start_day_offset,
                 "to": sh.end_time.strftime("%H:%M"), "to_day": sh.end_day_offset} for sh in v.shifts for r in roles]
    else:
        summ = summarize(v)
        s, e = summ["operating_start_min"], summ["operating_end_min"]
        fmt = lambda m: f"{(m % 1440) // 60:02d}:{m % 60:02d}"  # noqa: E731
        crew = [{"role": r, "shift": None, "from": fmt(s) if s is not None else None, "from_day": (s or 0) // 1440,
                 "to": fmt(e) if e is not None else None, "to_day": (e or 0) // 1440} for r in roles]
    reasons = []
    if cc:
        ts.ensure_reasons(db, cc.company_id)
        # the printed form asks for the reason *number* (digits read reliably when scanned); reason_id is stable
        reasons = [{"id": r.reason_id, "code": r.reason_code, "name": r.reason_name,
                    "name_local": r.reason_name_local} for r in db.query(TripReason).filter(
            TripReason.company_id == cc.company_id, TripReason.is_active.is_(True)).order_by(TripReason.reason_code)]
    return {"schedule_id": schedule.schedule_id, "schedule_code": schedule.schedule_code,
            "schedule_name": schedule.schedule_name, "version_number": v.version_number, "operating_date": on,
            "fleet_number": schedule.vehicle.fleet_number, "registration_number": schedule.vehicle.registration_number,
            "company_name": cc.company.company_name if cc else None,
            "cost_center": f"{cc.cost_center_code} · {cc.cost_center_name}" if cc else None,
            "paper_sheet_no": sheet.paper_sheet_no if sheet else None, "crew": crew,
            "version_valid_from": v.valid_from, "version_valid_to": v.valid_to,
            "runs": [ts.activity_planned(a) for a in sorted(v.activities, key=lambda x: x.sequence_number)],
            "reasons": reasons, "language": print_language(cc)}


# --------------------------------------------------------------------------
# Duty roster: planned against actual
# --------------------------------------------------------------------------
def _duty_label(x) -> tuple[str, str]:
    if x.duty_type != "SCHEDULE":
        return x.duty_type.title(), f"{x.duty_type.title()} {x.slot_index}"
    code = x.schedule.schedule_code if x.schedule else "?"
    return code, x.crew_role.title() + (f" - {x.shift.shift_name}" if x.shift else "")


def duty_day_summary(db: Session, cc_ids: list[int], on: dt.date) -> dict:
    """The end-of-day sheet: every duty slot, who was planned, who worked it, and why not."""
    from app.models.models import DutyActual, DutyRoster
    rosters = {r.roster_id: r for r in db.query(DutyRoster).filter(
        DutyRoster.cost_center_id.in_(cc_ids), DutyRoster.period_from <= on,
        DutyRoster.period_to >= on).all()}
    rows, counts = [], defaultdict(int)
    if rosters:
        for x in (db.query(DutyActual)
                  .filter(DutyActual.roster_id.in_(rosters), DutyActual.duty_date == on)
                  .order_by(DutyActual.duty_type, DutyActual.schedule_id, DutyActual.slot_index).all()):
            counts[x.status] += 1
            code, slot = _duty_label(x)
            rows.append({
                "roster_code": rosters[x.roster_id].roster_code,
                "cost_center": _cc_code(db, rosters[x.roster_id].cost_center_id),
                "duty_type": x.duty_type, "schedule_code": code,
                "fleet_number": x.schedule.vehicle.fleet_number if x.schedule else None,
                "slot_label": slot, "crew_role": x.crew_role,
                "segment": x.segment,
                "from_time": x.duty_start_time.strftime("%H:%M") if x.duty_start_time else None,
                "to_time": x.duty_end_time.strftime("%H:%M") if x.duty_end_time else None,
                "planned_name": x.planned_employee.full_name if x.planned_employee else None,
                "planned_number": x.planned_employee.employee_number if x.planned_employee else None,
                "actual_name": x.actual_employee.full_name if x.actual_employee else None,
                "actual_number": x.actual_employee.employee_number if x.actual_employee else None,
                "status": x.status, "reason": x.reason.reason_name if x.reason else None,
                "remarks": x.remarks})
    total = len(rows)
    changed = total - counts.get("AS_PLANNED", 0)
    return {"on": on, "slots": total, "counts": dict(counts), "changed": changed,
            "as_planned_pct": round(100 * counts.get("AS_PLANNED", 0) / total, 1) if total else None,
            "rows": rows}


def duty_statistics(db: Session, cc_ids: list[int], date_from: dt.date, date_to: dt.date) -> dict:
    """Per employee over a stretch of days: duties planned, duties worked, and the gap between.

    'Stood in' counts duties somebody worked that were planned for another person; 'missed'
    counts their own planned duties that somebody else worked or nobody did.
    """
    from app.models.models import DutyActual, DutyRoster
    check_range(date_from, date_to)
    roster_ids = [r.roster_id for r in db.query(DutyRoster).filter(
        DutyRoster.cost_center_id.in_(cc_ids), DutyRoster.period_from <= date_to,
        DutyRoster.period_to >= date_from).all()]
    rows = (db.query(DutyActual)
            .filter(DutyActual.roster_id.in_(roster_ids),
                    DutyActual.duty_date.between(date_from, date_to)).all()) if roster_ids else []

    people, totals, reasons = {}, defaultdict(int), defaultdict(int)

    def slot(employee) -> dict:
        key = employee.employee_id
        if key not in people:
            people[key] = {"employee_id": key, "employee_number": employee.employee_number,
                           "employee_name": employee.full_name, "staff_type": employee.staff_type,
                           "planned": 0, "worked": 0, "stood_in": 0, "missed": 0,
                           "spare": 0, "shunting": 0, "minutes": 0}
        return people[key]

    for x in rows:
        totals[x.status] += 1
        totals["slots"] += 1
        if x.reason is not None:
            reasons[x.reason.reason_name] += 1
        if x.planned_employee is not None:
            p = slot(x.planned_employee)
            p["planned"] += 1
            if x.actual_employee_id != x.planned_employee_id:
                p["missed"] += 1
        if x.actual_employee is not None:
            from app.services.duty_actual_service import _actual_span
            a = slot(x.actual_employee)
            a["worked"] += 1
            span = _actual_span(x)
            if span is not None:
                a["minutes"] += span[1] - span[0]
            if x.planned_employee_id and x.actual_employee_id != x.planned_employee_id:
                a["stood_in"] += 1
            if x.duty_type == "SPARE":
                a["spare"] += 1
            elif x.duty_type == "SHUNTING":
                a["shunting"] += 1

    items = sorted(people.values(), key=lambda p: (-p["worked"], p["employee_number"]))
    slots = totals.get("slots", 0)
    return {
        "date_from": date_from, "date_to": date_to, "slots": slots,
        "counts": {k: v for k, v in totals.items() if k != "slots"},
        "as_planned_pct": round(100 * totals.get("AS_PLANNED", 0) / slots, 1) if slots else None,
        "top_reasons": sorted(({"reason": k, "count": v} for k, v in reasons.items()),
                              key=lambda r: -r["count"])[:10],
        "crew": items,
    }


def duty_statistics_xlsx(rep: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Crew", ["Emp. no.", "Name", "Type", "Planned", "Worked", "Stood in", "Missed",
                        "Spare", "Shunting", "Hours"],
           [[c["employee_number"], c["employee_name"], c["staff_type"], c["planned"], c["worked"],
             c["stood_in"], c["missed"], c["spare"], c["shunting"],
             round(c["minutes"] / 60, 1) if c["minutes"] else 0] for c in rep["crew"]],
           [12, 26, 14, 10, 10, 10, 10, 10, 10, 10])
    _sheet(wb, "Reasons", ["Reason", "Times"],
           [[r["reason"], r["count"]] for r in rep["top_reasons"]], [40, 10])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def duty_day_xlsx(rep: dict) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    _sheet(wb, "Duties", ["Bus / duty", "Slot", "From", "To", "Planned", "Actual", "Status",
                          "Reason", "Remarks"],
           [[f"{r['fleet_number'] or r['duty_type'].title()} {r['schedule_code']}".strip(),
             r["slot_label"], r["from_time"] or "", r["to_time"] or "",
             r["planned_name"] or "", r["actual_name"] or "", r["status"],
             r["reason"] or "", r["remarks"] or ""] for r in rep["rows"]],
           [24, 20, 8, 8, 24, 24, 14, 24, 30])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
