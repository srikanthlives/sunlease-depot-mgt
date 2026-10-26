import datetime as dt
import io

import requests
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.enums import EmployeeSource, StaffType
from app.models.models import Employee, Location
from app.schemas.schemas import EmployeeIn

TEMPLATE_COLUMNS = ["Employee Number", "Full Name", "Designation", "Staff Type", "Home Depot Code",
                    "Start Date", "End Date", "Status"]
DATE_FORMAT = "%d-%b-%Y"


def parse_date(v) -> dt.date | None:
    """Accepts DD-MMM-YYYY (the app format), ISO dates and real Excel dates."""
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    s = str(v).strip()
    for fmt in (DATE_FORMAT, "%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognised date '{s}' (use DD-MMM-YYYY)")


def check_overlap(db: Session, number: str, start: dt.date, end: dt.date | None, exclude_id: int | None = None):
    """Stints of the same employee number must not overlap."""
    for o in db.query(Employee).filter(Employee.employee_number == number).all():
        if exclude_id and o.employee_id == exclude_id:
            continue
        o_end = o.end_date or dt.date.max
        if start <= o_end and (o.start_date <= (end or dt.date.max)):
            raise HTTPException(409, f"Employment dates overlap an existing stint of {number} "
                                     f"({o.start_date:%d-%b-%Y} to {o.end_date:%d-%b-%Y})" if o.end_date else
                                     f"Employment dates overlap an existing open stint of {number} (from {o.start_date:%d-%b-%Y})")


def template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Employees"
    ws.append(TEMPLATE_COLUMNS)
    ws.append(["E1001", "Ravi Kumar", "Driver", "DRIVER", "DEP01", "01-Apr-2025", "", "ACTIVE"])
    ws.append(["Staff Type values: " + ", ".join(StaffType.ALL)])
    for col, w in zip("ABCDEFGH", (16, 26, 18, 16, 16, 14, 14, 10)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upsert(db: Session, data: EmployeeIn, source: str, cost_center_id: int) -> str:
    """Upsert by (employee_number, start_date) into `cost_center_id`. Returns 'created' or 'updated'."""
    row = db.query(Employee).filter(Employee.employee_number == data.employee_number,
                                    Employee.start_date == data.start_date).first()
    if row is not None and row.cost_center_id != cost_center_id:
        raise HTTPException(409, f"Employee {data.employee_number} belongs to another cost center")
    check_overlap(db, data.employee_number, data.start_date, data.end_date, row.employee_id if row else None)
    if row is None:
        db.add(Employee(**data.model_dump(), source=source, last_synced_at=dt.datetime.utcnow(),
                        cost_center_id=cost_center_id))
        db.flush()
        return "created"
    for k, v in data.model_dump().items():
        setattr(row, k, v)
    row.source = source
    row.last_synced_at = dt.datetime.utcnow()
    db.flush()
    return "updated"


def _depots(db: Session, company_id: int) -> dict:
    return {l.location_code.upper(): l.location_id for l in db.query(Location).filter(Location.company_id == company_id)}


def import_workbook(db: Session, content: bytes, cost_center_id: int, company_id: int) -> dict:
    try:
        ws = load_workbook(io.BytesIO(content), data_only=True).active
    except Exception:
        raise HTTPException(400, "Could not read the file - upload an .xlsx workbook")
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "The workbook is empty")
    header = [str(h or "").strip().lower() for h in rows[0]]
    need = [c.lower() for c in TEMPLATE_COLUMNS]
    if any(c not in header for c in need[:2] + [need[5]]):
        raise HTTPException(400, f"Header row must include: {', '.join(TEMPLATE_COLUMNS)}")
    idx = {name: header.index(name) if name in header else None for name in need}
    depots = _depots(db, company_id)
    created = updated = 0
    errors = []
    parsed = []
    for n, r in enumerate(rows[1:], start=2):
        def cell(name):
            i = idx[name.lower()]
            return r[i] if i is not None and i < len(r) else None
        if not any(r) or cell("Full Name") in (None, "") and cell("Employee Number") in (None, ""):
            continue
        try:
            depot_code = str(cell("Home Depot Code") or "").strip().upper()
            if depot_code and depot_code not in depots:
                raise ValueError(f"unknown depot code '{depot_code}'")
            parsed.append(EmployeeIn(
                employee_number=str(cell("Employee Number") or "").strip(),
                full_name=str(cell("Full Name") or "").strip(),
                designation=(str(cell("Designation")).strip() if cell("Designation") else None),
                staff_type=str(cell("Staff Type") or StaffType.DRIVER).strip().upper(),
                home_depot_id=depots.get(depot_code) if depot_code else None,
                start_date=parse_date(cell("Start Date")),
                end_date=parse_date(cell("End Date")),
                status=str(cell("Status") or "ACTIVE").strip().upper()))
        except Exception as exc:
            errors.append({"row": n, "error": str(exc).split("\n")[0][:200]})
    if errors:  # all-or-nothing: nothing is saved if any row is bad
        return {"created": 0, "updated": 0, "errors": errors}
    try:
        for data in parsed:
            if _upsert(db, data, EmployeeSource.UPLOAD, cost_center_id) == "created":
                created += 1
            else:
                updated += 1
    except HTTPException as exc:
        db.rollback()
        return {"created": 0, "updated": 0, "errors": [{"row": None, "error": exc.detail}]}
    return {"created": created, "updated": updated, "errors": []}


def sync_from_hrms(db: Session, cost_center_id: int, company_id: int) -> dict:
    """Pull employees from the HRMS export endpoint. Expected JSON: a list of
    {employee_number, full_name, designation, staff_type, home_depot_code,
     start_date, end_date, status} (ISO or DD-MMM-YYYY dates)."""
    if not settings.HRMS_SYNC_URL:
        raise HTTPException(400, "HRMS sync is not configured (set DEPOT_HRMS_SYNC_URL and DEPOT_HRMS_SYNC_TOKEN)")
    try:
        headers = {"Authorization": f"Bearer {settings.HRMS_SYNC_TOKEN}"} if settings.HRMS_SYNC_TOKEN else {}
        resp = requests.get(settings.HRMS_SYNC_URL, headers=headers, timeout=30)
        resp.raise_for_status()
        payload = resp.json()
    except Exception as exc:
        raise HTTPException(502, f"Could not fetch employees from HRMS: {exc}")
    depots = _depots(db, company_id)
    created = updated = 0
    errors = []
    for n, item in enumerate(payload if isinstance(payload, list) else payload.get("employees", []), start=1):
        try:
            code = str(item.get("home_depot_code") or "").strip().upper()
            data = EmployeeIn(
                employee_number=str(item["employee_number"]).strip(), full_name=item["full_name"],
                designation=item.get("designation"),
                staff_type=str(item.get("staff_type") or StaffType.DRIVER).upper(),
                home_depot_id=depots.get(code), start_date=parse_date(item["start_date"]),
                end_date=parse_date(item.get("end_date")), status=str(item.get("status") or "ACTIVE").upper())
            if _upsert(db, data, EmployeeSource.HRMS, cost_center_id) == "created":
                created += 1
            else:
                updated += 1
        except Exception as exc:
            detail = exc.detail if isinstance(exc, HTTPException) else str(exc).split("\n")[0]
            errors.append({"row": n, "error": str(detail)[:200]})
    return {"created": created, "updated": updated, "errors": errors}


def references(db: Session, employee_id: int) -> dict:
    """Where an employee is already used. A used stint is never moved between cost centers:
    the history of a depot has to stay with that depot."""
    from app.models.models import DutyActual, DutyAssignment, DutyPatternMember, TripSheetCrew
    counts = {
        "trip_sheets": db.query(TripSheetCrew).filter(TripSheetCrew.employee_id == employee_id).count(),
        "duty_assignments": db.query(DutyAssignment).filter(DutyAssignment.employee_id == employee_id).count(),
        "duty_actuals": db.query(DutyActual).filter(
            (DutyActual.planned_employee_id == employee_id)
            | (DutyActual.actual_employee_id == employee_id)).count(),
        "duty_patterns": db.query(DutyPatternMember).filter(
            DutyPatternMember.employee_id == employee_id).count(),
    }
    return {k: v for k, v in counts.items() if v}


def transfer(db: Session, e: Employee, target_cc: int, effective_from: dt.date | None, user) -> dict:
    """Move an employee to another cost center.

    Two different things wear the same name, so both are handled here:

    * an employee entered against the wrong depot and not yet used anywhere is a **correction**:
      the row simply moves.
    * an employee who has already worked here is a **transfer**: their stint at this depot ends and
      a new stint starts at the other one, which is what the per-stint `employees` table is for.
      Their old duties and trip sheets stay with the depot where they happened.
    """
    if e.cost_center_id == target_cc:
        raise HTTPException(409, "The employee is already in that cost center")
    used = references(db, e.employee_id)
    if not used:
        e.cost_center_id = target_cc
        # the home depot is company-level, so it stays valid; the caller has checked the company
        return {"action": "MOVED", "employee_id": e.employee_id, "used": {}}

    if effective_from is None:
        raise HTTPException(422, "This employee has already worked in this cost center, so the "
                                 "transfer needs a date: their stint here ends the day before it")
    if effective_from <= e.start_date:
        raise HTTPException(422, f"The transfer date must be after the stint started "
                                 f"({e.start_date:%d-%b-%Y})")
    if e.end_date and effective_from > e.end_date:
        raise HTTPException(422, f"The stint here already ended on {e.end_date:%d-%b-%Y}")
    last_day = effective_from - dt.timedelta(days=1)
    check_overlap(db, e.employee_number, effective_from, None, exclude_id=e.employee_id)
    e.end_date = last_day
    new = Employee(employee_number=e.employee_number, full_name=e.full_name, designation=e.designation,
                   staff_type=e.staff_type, home_depot_id=None, start_date=effective_from,
                   status=e.status, source=e.source, cost_center_id=target_cc)
    db.add(new)
    db.flush()
    return {"action": "TRANSFERRED", "employee_id": new.employee_id, "ended_on": last_day,
            "used": used}
