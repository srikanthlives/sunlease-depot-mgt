"""Billing (approved sheets only), missing sheets, crew hours and the printable sheet."""
import io

from openpyxl import load_workbook

from .conftest import login
from .test_trip_sheets import as_planned, driver_on, new_sheet, published, reason, save, staff


def approve(world, sheet):
    c, h = world["client"], world["h"]
    assert c.post(f"/api/trip-sheets/{sheet['sheet_id']}/submit", headers=h).status_code == 200
    assert c.post(f"/api/trip-sheets/{sheet['sheet_id']}/approve", headers=h).status_code == 200


def test_billing_counts_only_approved_sheets_and_splits_losses(world):
    s, _ = published(world)
    c, h = world["client"], world["h"]
    d = staff(world, "D1")
    # day 1: T1 cancelled (external, deducted), T2 cut short at MID (client, still billed)
    day1 = new_sheet(world, s, "2026-11-05")
    down = next(r for r in day1["rows"] if r["planned"].get("trip_number") == "T2")["planned"]
    mid = next(st for st in down["stops"] if st["location_code"] == "MID")
    rows = as_planned(day1, T1={"status": "CANCELLED", "reason_id": reason(world, "ROAD_CLOSED"),
                                "start_odometer_km": None, "end_odometer_km": None},
                      T2={"status": "PARTIAL", "reason_id": reason(world, "CLIENT"), "stopped_at_route_stop_id": mid["route_stop_id"]})
    save(world, day1, rows, driver_on(day1, d))
    approve(world, day1)
    # day 2: as planned but only submitted - not billed yet
    day2 = new_sheet(world, s, "2026-11-06")
    save(world, day2, as_planned(day2, odo=1030.5), driver_on(day2, d))   # opens at day 1's closing reading
    assert c.post(f"/api/trip-sheets/{day2['sheet_id']}/submit", headers=h).status_code == 200
    rep = c.get("/api/reports/billing?date_from=2026-11-01&date_to=2026-11-30", headers=h).json()
    assert rep["not_approved_sheets"] == 1 and len(rep["daily"]) == 1
    g = rep["rows"][0]
    assert g["days"] == 1 and g["planned_billable_km"] == 20.5
    assert g["lost_external_km"] == 10 and g["lost_client_km"] == 4 and g["lost_operator_km"] == 0
    assert g["deducted_km"] == 10 and g["lost_billed_km"] == 4 and g["billable_km"] == 10.5
    assert rep["totals"]["billable_km"] == 10.5
    assert {x["status"] for x in rep["losses"]} == {"CANCELLED", "PARTIAL"}
    # export: three tabs, totals line
    r = c.get("/api/reports/billing.xlsx?date_from=2026-11-01&date_to=2026-11-30", headers=h)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Daily", "Losses"]
    summary = list(wb["Summary"].values)
    assert summary[-1][0] == "TOTAL" and summary[-1][10] == 10.5
    # export needs a publisher; viewers can still see the report
    viewer = login(c, "viewer", "Viewer@123")
    assert c.get("/api/reports/billing?date_from=2026-11-01&date_to=2026-11-30", headers=viewer).status_code == 200
    assert c.get("/api/reports/billing.xlsx?date_from=2026-11-01&date_to=2026-11-30", headers=viewer).status_code == 403
    # ranges are bounded
    assert c.get("/api/reports/billing?date_from=2026-01-01&date_to=2026-12-31", headers=h).status_code == 422


def test_missing_sheets(world):
    s, _ = published(world)
    c, h = world["client"], world["h"]
    sheet = new_sheet(world, s, "2026-11-02")
    save(world, sheet, as_planned(sheet), driver_on(sheet, staff(world, "D1")))
    approve(world, sheet)
    new_sheet(world, s, "2026-11-03")                                   # draft
    rep = c.get("/api/reports/missing-sheets?date_from=2026-10-30&date_to=2026-11-04", headers=h).json()
    assert rep["expected"] == 4                                          # effective from 2026-11-01
    assert rep["counts"] == {"APPROVED": 1, "DRAFT": 1, "NOT_STARTED": 2}
    assert [(r["operating_date"], r["status"]) for r in rep["rows"]] == [
        ("2026-11-01", "NOT_STARTED"), ("2026-11-03", "DRAFT"), ("2026-11-04", "NOT_STARTED")]


def test_crew_hours_with_handover(world):
    s, _ = published(world)
    c, h = world["client"], world["h"]
    a, b = staff(world, "D1"), staff(world, "D2")
    sheet = new_sheet(world, s)
    sheet = save(world, sheet, as_planned(sheet), driver_on(sheet, a)).json()
    c.post(f"/api/trip-sheet-crew/{sheet['crew'][0]['crew_id']}/handover", headers=h,
           json={"at_time": "07:00", "employee_id": b, "reason": "unwell"})
    rep = c.get("/api/reports/crew-hours?date_from=2026-11-01&date_to=2026-11-30", headers=h).json()
    hours = {e["employee_number"]: e["minutes"] for e in rep["employees"]}
    assert hours == {"D1": 120, "D2": 105} and rep["total_minutes"] == 225
    took = next(x for e in rep["employees"] for x in e["items"] if e["employee_number"] == "D2")
    assert took["took_over"] and took["handover_reason"] == "unwell"
    r = c.get("/api/reports/crew-hours.xlsx?date_from=2026-11-01&date_to=2026-11-30", headers=h)
    assert load_workbook(io.BytesIO(r.content)).sheetnames == ["Summary", "Duties"]


def test_print_data_follows_the_published_plan(world):
    s, _ = published(world)
    c, h = world["client"], world["h"]
    d = c.get(f"/api/trip-sheets/print-data?date=2026-11-05&schedule_id={s['schedule_id']}", headers=h).json()[0]
    assert len(d["runs"]) == 5 and d["runs"][1]["trip_number"] == "T1"
    assert [x["role"] for x in d["crew"]] == ["DRIVER"] and d["crew"][0]["from"] == "05:00"
    assert d["language"] == "te" and any(r["code"] == "BREAKDOWN" and r["id"] for r in d["reasons"])
    # the language belongs to the cost center, and to the company when the depot has no choice of its own
    cc = c.get("/api/cost-centers", headers=h).json()[0]
    assert c.put(f"/api/cost-centers/{cc['cost_center_id']}", headers=h,
                 json={**cc, "print_language": "kn"}).status_code == 200
    assert c.get("/api/trip-sheets/print-data?date=2026-11-05", headers=h).json()[0]["language"] == "kn"
    assert c.put(f"/api/cost-centers/{cc['cost_center_id']}", headers=h,
                 json={**cc, "print_language": "fr"}).status_code == 422
    # the language belongs to the schedule's own cost center, not to the cost center the request is made in,
    # so "Print all sheets" from the combined view prints each bus in its own depot's language
    everything = {**h, "X-Cost-Center-Id": "all"}
    assert c.get("/api/trip-sheets/print-data?date=2026-11-05", headers=everything).json()[0]["language"] == "kn"
    # clearing the depot's own choice falls back to the company's
    assert c.put(f"/api/cost-centers/{cc['cost_center_id']}", headers=h,
                 json={**cc, "print_language": None}).status_code == 200
    company = c.get("/api/companies", headers=h).json()[0]
    assert c.put(f"/api/companies/{company['company_id']}", headers=h,
                 json={**company, "print_language": "ta"}).status_code == 200
    assert c.get("/api/trip-sheets/print-data?date=2026-11-05", headers=h).json()[0]["language"] == "ta"
    assert c.get(f"/api/trip-sheets/print-data?date=2026-10-01&schedule_id={s['schedule_id']}", headers=h).status_code == 404


def test_reason_codes_are_configurable_and_print_with_the_regional_name(world):
    """Depots add their own fault reasons; the printed legend carries the regional name in brackets."""
    s, _ = published(world)
    c, h = world["client"], world["h"]
    # the nine defaults exist on first use, and can be renamed in the regional language
    seeded = c.get("/api/trip-reasons", headers=h).json()
    assert {r["reason_code"] for r in seeded} >= {"BREAKDOWN", "CLIENT", "OTHER"}
    bd = next(r for r in seeded if r["reason_code"] == "BREAKDOWN")
    assert c.put(f"/api/trip-reasons/{bd['reason_id']}", headers=h,
                 json={**bd, "reason_name_local": "వాహనం మొరాయించింది"}).status_code == 200
    # a depot-specific reason
    new = c.post("/api/trip-reasons", headers=h, json={"reason_code": "TYRE", "reason_name": "Tyre puncture",
                                                       "reason_name_local": "టైర్ పంక్చర్", "default_attributable_to": "OPERATOR"})
    assert new.status_code == 201
    tyre = new.json()
    assert c.post("/api/trip-reasons", headers=h, json={"reason_code": "TYRE", "reason_name": "x"}).status_code == 409
    # both reach the printed sheet, keyed by the number the driver writes
    d = c.get(f"/api/trip-sheets/print-data?date=2026-11-05&schedule_id={s['schedule_id']}", headers=h).json()[0]
    printed = {r["code"]: r for r in d["reasons"]}
    assert printed["BREAKDOWN"]["name_local"] == "వాహనం మొరాయించింది"
    assert printed["TYRE"]["name"] == "Tyre puncture" and printed["TYRE"]["id"] == tyre["reason_id"]
    # an unused reason can go; one written on a sheet cannot
    assert c.delete(f"/api/trip-reasons/{tyre['reason_id']}", headers=h).status_code == 204
    sheet = new_sheet(world, s, "2026-11-07")
    save(world, sheet, as_planned(sheet, T1={"status": "CANCELLED", "reason_id": bd["reason_id"],
                                             "start_odometer_km": None, "end_odometer_km": None}),
         driver_on(sheet, staff(world, "D9")))
    assert c.delete(f"/api/trip-reasons/{bd['reason_id']}", headers=h).status_code == 409
    # deactivating keeps it on the sheets that use it
    assert c.put(f"/api/trip-reasons/{bd['reason_id']}", headers=h, json={**bd, "is_active": False}).status_code == 200
    assert all(r["code"] != "BREAKDOWN" for r in c.get("/api/trip-sheets/print-data?date=2026-11-05",
                                                       headers=h).json()[0]["reasons"])
    # a scheduler may configure them; a viewer may not
    viewer = login(c, "viewer", "Viewer@123")
    assert c.post("/api/trip-reasons", headers=viewer, json={"reason_code": "Z", "reason_name": "z"}).status_code == 403
